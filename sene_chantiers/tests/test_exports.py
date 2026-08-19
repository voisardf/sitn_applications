"""Exports PDF et Excel."""

import base64
import io
from datetime import timedelta
from unittest.mock import patch

from django.core.files.base import ContentFile
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from ..labels import appreciation_scale
from ..models import (
    Appreciation,
    EmailRecord,
    EmailStatus,
    EmailTemplate,
    Photo,
    PhotoStatus,
)
from ..services import excel, pdf
from .factories import (
    TODAY,
    jpeg_bytes,
    make_chantier,
    make_control_report,
    make_followup,
    make_measure,
    make_user,
    seed_sections,
)


class MemberClientMixin:
    def setUp(self):
        super().setUp()
        self.user = make_user("membre")
        self.client = Client()
        self.client.force_login(self.user)


class PdfExportTest(MemberClientMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.chantier = make_chantier(satac_number=999700)
        self.report = make_control_report(chantier=self.chantier)
        seed_sections(self.report)
        make_measure(self.report)

    def test_the_first_page_is_aired_and_section_04_fills_its_page(self):
        """Mise en page : sections 01–02 aérées, section 04 occupant sa page.

        Les valeurs exactes sont un choix de maquette ; ce qui compte et
        se casse silencieusement, c'est le rapport entre elles — un en-tête
        plus serré que le corps, ou une section 04 tassée alors qu'elle a
        une page pour elle seule.
        """
        html = pdf.render_html(
            "sene_chantiers/pdf/control_report.html",
            pdf._control_report_context(self.report),
        )

        def millimetres(rule, prop):
            block = html.split(rule, 1)[1].split("}", 1)[0]
            value = block.split(prop, 1)[1].split(";", 1)[0]
            first = value.split(":")[-1].strip().split()[0]
            return float(first.replace("mm", ""))

        # L'espace sous le bandeau d'en-tête dépasse celui entre sections :
        # c'est ce qui aère le haut de la première page.
        between_sections = millimetres(".section {", "margin-bottom")
        under_header = millimetres(".metastrip {", "margin-bottom")
        self.assertGreater(under_header, between_sections)

        # La section 04 a sa page : ses lignes respirent au lieu d'être
        # compressées pour gagner une place dont elle n'a pas besoin.
        compact = millimetres(".grid.compact td {", "padding")
        self.assertGreaterEqual(compact, 1.5)

    def test_html_carries_no_photograph(self):
        """Les photos vont au courriel, jamais dans le rapport."""
        photo = Photo(
            control_report=self.report,
            status=PhotoStatus.DONE,
            caption="Légende de la photo",
        )
        photo.image.save("p.jpg", ContentFile(jpeg_bytes()), save=True)

        html = pdf.render_html(
            "sene_chantiers/pdf/control_report.html",
            pdf._control_report_context(self.report),
        )
        self.assertNotIn(photo.image.name, html)
        self.assertNotIn("Légende de la photo", html)
        self.assertNotIn("file://", html)
        self.assertIn("Rapport de contrôle environnemental", html)
        self.assertIn("Mesures à prendre", html)

    def test_the_only_embedded_images_are_the_logo_and_watermark(self):
        """Le gabarit embarque le logo et le filigrane, rien d'autre.

        Une photo embarquée se verrait ici : le service ne résout aucune
        URL, donc toute image du PDF est forcément une data URI.
        """
        photo = Photo(control_report=self.report, status=PhotoStatus.DONE)
        photo.image.save("p.jpg", ContentFile(jpeg_bytes()), save=True)

        html = pdf.render_html(
            "sene_chantiers/pdf/control_report.html",
            pdf._control_report_context(self.report),
        )
        self.assertEqual(html.count("data:image"), 2)
        self.assertIn("data:image/png;base64,", html)      # logo
        self.assertIn("data:image/svg+xml;base64,", html)  # filigrane

    def test_prototype_watermark_is_present_on_every_page(self):
        """Temporaire : à retirer avec le bandeau web à la mise en service."""
        html = pdf.render_html(
            "sene_chantiers/pdf/control_report.html",
            pdf._control_report_context(self.report),
        )
        # Posé en fond de @page, donc répété sur chaque page sans avoir à
        # l'insérer dans le flux.
        self.assertIn("background-image", html)
        self.assertIn("PROTOTYPE", base64.b64decode(
            pdf.watermark_data_uri().split("base64,", 1)[1]).decode())

    def test_logo_is_inlined_rather_than_linked(self):
        # A {% static %} URL would render as a silently missing image.
        self.assertTrue(pdf.logo_data_uri().startswith("data:image/png;base64,"))

    def _control_html(self):
        return pdf.render_html(
            "sene_chantiers/pdf/control_report.html",
            pdf._control_report_context(self.report),
        )

    def test_measures_checklist_and_signature_each_start_a_page(self):
        """Three deliberate breaks: after 02, after 03, after 04."""
        self.assertEqual(self._control_html().count('break-after"'), 3)

    def test_a_compliant_first_control_keeps_its_short_sections_together(self):
        """With nothing to prescribe, section 03 must not claim a page.

        The break after 02 is conditional for exactly this case: a
        compliant control would otherwise print a page holding one line.
        """
        self.report.corrective_measures.all().delete()
        html = self._control_html()
        self.assertEqual(html.count('break-after"'), 2)
        self.assertIn("Aucune mesure à prendre", html)

    def test_followup_splits_after_the_measures_table(self):
        followup = make_followup(self.chantier)
        html = pdf.render_html(
            "sene_chantiers/pdf/corrective_measure_report.html",
            {
                "report": followup,
                "chantier": self.chantier,
                "first_control_date": self.report.control_date,
                "observations": [],
                "report_date": followup.follow_up_date,
                "appreciation_scale": appreciation_scale("followup"),
                "conclusion_text": "",
                "logo_data_uri": "",
            },
        )
        self.assertEqual(html.count('break-after"'), 1)

    def test_appreciation_card_is_tinted_with_the_level_in_force(self):
        self.report.global_appreciation = Appreciation.ROUGE
        self.report.save()
        self.assertIn("appr-card tint-rouge", self._control_html())

    def test_print_rules_are_present(self):
        html = pdf.render_html(
            "sene_chantiers/pdf/control_report.html",
            pdf._control_report_context(self.report),
        )
        self.assertIn("break-inside: avoid", html)
        self.assertIn("counter(page)", html)

    @patch("sene_chantiers.services.pdf.requests.post")
    def test_service_response_is_returned(self, mock_post):
        mock_post.return_value.content = b"%PDF-1.7 ..."
        mock_post.return_value.raise_for_status = lambda: None
        response = self.client.get(
            reverse(
                "sene_chantiers:pdf_export", args=["controle", self.report.pk]
            )
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")

    @patch("sene_chantiers.services.pdf.requests.post")
    def test_non_pdf_response_is_rejected(self, mock_post):
        mock_post.return_value.content = b"<html>erreur</html>"
        mock_post.return_value.raise_for_status = lambda: None
        with self.assertRaises(pdf.PdfServiceError):
            pdf.report_pdf(self.report)

    @patch(
        "sene_chantiers.services.pdf.requests.post",
        side_effect=pdf.requests.RequestException("boom"),
    )
    def test_unreachable_service_degrades_gracefully(self, _post):
        response = self.client.get(
            reverse(
                "sene_chantiers:pdf_export", args=["controle", self.report.pk]
            ),
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "indisponible")


class ExportFilenameTest(MemberClientMixin, TestCase):
    """Un rapport, un nom — le même dans les trois sorties.

    Le n° de suivi manquait partout : trois rapports d'un même dossier
    arrivaient dans le dossier de téléchargements sous des noms
    indiscernables.
    """

    def setUp(self):
        super().setUp()
        self.chantier = make_chantier(satac_number=999900)
        self.report = make_control_report(chantier=self.chantier)
        seed_sections(self.report)
        self.followup = make_followup(self.chantier, user=self.user)

    @patch("sene_chantiers.services.pdf.html_to_pdf", return_value=b"%PDF-1.7")
    def test_the_pdf_filename_carries_the_followup_number(self, _render):
        cases = [
            (["controle", self.report.pk], "rapport_controle_999900.pdf"),
            (["suivi", self.followup.pk], "rapport_suivi1_999900.pdf"),
        ]
        for args, expected in cases:
            with self.subTest(expected=expected):
                response = self.client.get(
                    reverse("sene_chantiers:pdf_export", args=args))
                self.assertEqual(response.status_code, 200)
                self.assertIn(expected, response["Content-Disposition"])

    def test_the_email_attachment_uses_the_same_name(self):
        from django.core import mail

        from ..services import emails as email_service

        record = EmailRecord.objects.create(
            chantier=self.chantier,
            corrective_measure_report=self.followup,
            template_used=EmailTemplate.NON_CONFORMITES,
            recipient_email="destinataire@example.ch",
            subject="Objet",
            body="Corps",
        )
        email_service.send(record, pdf_bytes=b"%PDF-1.7")
        attachments = dict(
            (name, None) for name, _content, _type in mail.outbox[0].attachments
        )
        self.assertIn("rapport_suivi1_999900.pdf", attachments)

    def test_the_three_outputs_agree(self):
        """Aucune sortie ne doit inventer sa propre convention."""
        from ..labels import export_stem

        self.assertEqual(export_stem(self.report), "controle_999900")
        self.assertEqual(export_stem(self.followup), "suivi1_999900")


class ExcelExportTest(MemberClientMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.chantier = make_chantier(satac_number=999800)
        self.report = make_control_report(chantier=self.chantier)
        seed_sections(self.report)
        make_measure(self.report)
        make_followup(self.chantier)

    def test_workbook_structure(self):
        from openpyxl import load_workbook

        content = excel.dossier_workbook(self.chantier)
        workbook = load_workbook(io.BytesIO(content))
        self.assertEqual(
            workbook.sheetnames,
            [
                "Informations générales",
                "Contrôle initial",
                "Suivi 1",
                "Courriels envoyés",
            ],
        )
        general = workbook["Informations générales"]
        # Le n° SATAC reste numérique dans le classeur, pas converti en texte.
        self.assertEqual(general["B1"].value, self.chantier.satac_number)

    def test_view_streams_a_workbook(self):
        response = self.client.get(
            reverse("sene_chantiers:excel_export",
                    args=["controle", self.report.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("dossier_controle_999800.xlsx",
                      response["Content-Disposition"])

    def test_the_filename_names_the_report_it_was_launched_from(self):
        """Sinon chaque ligne télécharge le même nom de fichier."""
        # Créé après, mais daté avant : le numéro suit les dates, pas
        # l'ordre d'insertion.
        earlier = make_followup(self.chantier, user=self.user,
                                follow_up_date=TODAY - timedelta(days=10))
        later = self.chantier.corrective_measure_reports.exclude(
            pk=earlier.pk).get()
        cases = [
            (["controle", self.report.pk], "dossier_controle_999800.xlsx"),
            (["suivi", earlier.pk], "dossier_suivi1_999800.xlsx"),
            (["suivi", later.pk], "dossier_suivi2_999800.xlsx"),
        ]
        for args, expected in cases:
            with self.subTest(expected=expected):
                response = self.client.get(
                    reverse("sene_chantiers:excel_export", args=args))
                self.assertIn(expected, response["Content-Disposition"])

    def test_lazy_labels_are_resolved(self):
        """openpyxl refuse les proxies de traduction."""
        from openpyxl import load_workbook

        EmailRecord.objects.create(
            chantier=self.chantier,
            control_report=self.report,
            template_used=EmailTemplate.NON_CONFORMITES,
            recipient_email="destinataire@example.ch",
            subject="Objet",
            body="Corps",
            status=EmailStatus.SENT,
            sent_at=timezone.now(),
        )
        workbook = load_workbook(io.BytesIO(excel.dossier_workbook(self.chantier)))
        sheet = workbook["Courriels envoyés"]
        self.assertEqual(sheet.cell(row=2, column=2).value, "Mesures à prendre")


# ---------------------------------------------------------------------------
# Échéances
# ---------------------------------------------------------------------------
